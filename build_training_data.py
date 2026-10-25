"""Bangun dataset latih GABUNGAN dari 3 generasi form + Rekap_All, dan cek apakah datanya layak dipakai.

  python build_training_data.py --v1 Form_Loker_Tele_Marketing_Jawaban.xlsx --v2 Kuis_tanpa_judul.xlsx \\
                                --v3 Form_Data_Diri_Kandidat.xlsx --rekap Rekap_All.xlsx --out training_data

Keluaran di folder --out (SEMUA ada di .gitignore, jangan di-commit: berisi data pelamar):
  unified_all.csv         lengkap, ADA nama/HP/KTP (untuk pengecekan manual)
  unified_anonymized.csv  tanpa identitas -> aman dibagikan / dipakai di Colab
  laporan.md              cakupan label per bulan, kualitas jarak, dan hasil uji kelayakan model
"""
import argparse
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))
warnings.filterwarnings("ignore")

from config import SEED, UNIFIED_WINDOWS  # noqa: E402
from harmonize import CORE, build_unified, mark_training  # noqa: E402

NUM = ["usia", "jarak_km", "exp_tele_ord", "exp_finance_ord"]
BIN = ["fresh_grad", "exp_tele", "exp_finance", "lokasi_presisi"]
CAT = ["pend_level"]


def _pipe(num=NUM, bin_=BIN, cat=CAT):
    from sklearn.compose import ColumnTransformer
    from sklearn.impute import SimpleImputer
    from sklearn.linear_model import LogisticRegression
    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import OneHotEncoder, StandardScaler

    parts = [("n", Pipeline([("i", SimpleImputer(strategy="median")), ("s", StandardScaler())]), num),
             ("b", SimpleImputer(strategy="constant", fill_value=0), bin_)]
    if cat:
        parts.append(("c", OneHotEncoder(handle_unknown="ignore", min_frequency=8), cat))
    return Pipeline([("p", ColumnTransformer(parts)), ("m", LogisticRegression(C=0.3, class_weight="balanced", max_iter=2000))])


def _cv_auc(d, **kw):
    from sklearn.metrics import roc_auc_score
    from sklearn.model_selection import RepeatedStratifiedKFold

    d = d.reset_index(drop=True)
    a = []
    for tr, te in RepeatedStratifiedKFold(n_splits=5, n_repeats=5, random_state=SEED).split(d, d["y"]):
        m = _pipe(**kw).fit(d.iloc[tr], d["y"].iloc[tr])
        a.append(roc_auc_score(d["y"].iloc[te], m.predict_proba(d.iloc[te])[:, 1]))
    return float(np.mean(a))


def _transfer(tr, te):
    from sklearn.metrics import roc_auc_score

    m = _pipe().fit(tr, tr["y"])
    return float(roc_auc_score(te["y"], m.predict_proba(te)[:, 1]))


def evaluate(u: pd.DataFrame) -> list[str]:
    t = u[u["use_train"]].copy()
    t["jarak_km"] = t["jarak_km"].clip(upper=100)
    t["win"] = np.where((t["era"] == "v1") & (t["ts"] < "2025-01-01"), "v1_2024", np.where(t["era"] == "v1", "v1_2026", "v2_2026"))
    L = ["| Jendela data | n | % lolos | AUC CV (fitur inti) |", "|---|---|---|---|"]
    for w, d in t.groupby("win"):
        L.append(f"| {w} | {len(d)} | {d['y'].mean():.0%} | {_cv_auc(d):.3f} |")
    L.append(f"| **gabungan semua** | {len(t)} | {t['y'].mean():.0%} | {_cv_auc(t):.3f} |")
    L += ["", "Uji transfer (latih di satu periode, uji di periode lain; AUC 0,5 = tebakan acak):", "",
          "| Latih | Uji | AUC |", "|---|---|---|"]
    g = {w: d for w, d in t.groupby("win")}
    for a, b in (("v1_2024", "v2_2026"), ("v1_2024", "v1_2026"), ("v2_2026", "v1_2024"), ("v2_2026", "v1_2026")):
        if a in g and b in g:
            L.append(f"| {a} | {b} | {_transfer(g[a], g[b]):.3f} |")
    return L


def report(u: pd.DataFrame, eval_lines: list[str]) -> str:
    u = u.copy()
    u["bulan"] = u["ts"].dt.to_period("M").astype(str)
    cov = u.groupby(["era", "bulan"]).agg(n=("y", "size"), lolos=("y", "sum"), tidak_jelas=("y", lambda s: int(s.isna().sum())),
                                           pct=("y", "mean"), dipakai=("use_train", "sum"))
    L = ["# Laporan dataset gabungan (Tele Marketing)", "", f"Pelamar unik: **{len(u)}** · dipakai untuk training: **{int(u['use_train'].sum())}**", "",
         "## Cakupan label per bulan submit", "", "| Era | Bulan | n | lolos | label tidak jelas | % lolos | dipakai training |", "|---|---|---|---|---|---|---|"]
    for (era, b), r in cov.iterrows():
        L.append(f"| {era} | {b} | {int(r.n)} | {int(r.lolos)} | {int(r.tidak_jelas)} | {r.pct:.0%} | {int(r.dipakai)} |")
    L += ["", "Bulan dengan % lolos ≈ 0 (mis. 2024-12 s/d 2025-11) = Rekap_All tidak memuat bulan itu, BUKAN berarti semua ditolak -> dikeluarkan.",
          "Bulan terakhir (mis. 2026-09/10) = pelamar belum sempat diproses saat rekap dibuat -> dikeluarkan.", "",
          "## Kualitas jarak (sumber nilai jarak_km)", ""]
    q = u.groupby(["era", "jarak_src"]).size().unstack(fill_value=0)
    L += [q.to_markdown(), "", "bracket = alamat bebas tak terbaca/menyimpang, dipakai titik tengah pilihan jarak yang diisi pelamar sendiri.", "",
          "## Uji kelayakan model (fitur inti, regresi logistik)", ""] + eval_lines
    return "\n".join(L)


KELOMPOK = {"A": "A. 2024 (form lama)", "B": "B. Des 2025 - Jan 2026 (form lama)", "C": "C. Kuis 2026", "D": "D. Data Diri 2026 (Tele Marketing)"}
COLS = [  # (kolom di u, nama kolom Excel, penjelasan)
    ("id", "ID", "Nomor urut baris"), ("kelompok", "Kelompok Data", "A-D = asal & periode data (lihat sheet Ringkasan)"),
    ("era", "Sumber Form", "v1 = Form Loker Tele Marketing; v2 = Kuis tanpa judul; v3 = Form Data Diri Kandidat"),
    ("ts", "Tanggal Submit", "Waktu pelamar mengisi form (submit terakhir sebelum interview)"),
    ("nama", "Nama", "IDENTITAS - jangan dipakai sebagai fitur"), ("hp", "No HP", "IDENTITAS - kunci pencocokan ke rekap"),
    ("ktp", "NIK/KTP", "IDENTITAS - kunci data unik"),
    ("y", "LABEL (1=lanjut interview)", "1 = ada di Rekap_All (lolos seleksi berkas); 0 = tidak ada di rekap; kosong = belum diketahui"),
    ("label_note", "Keterangan Label", "Alasan label"), ("rekap_date", "Tanggal Interview (Rekap)", "Tanggal interview di Rekap_All"),
    ("siap", "Siap Dipakai Training?", "Ya / Hati-hati / Tidak (alasan di kolom sebelahnya)"), ("siap_alasan", "Alasan", ""),
    ("usia", "Usia", "tahun, dari tanggal lahir (cadangan: titik tengah pilihan usia di form kuis)"),
    ("pend_level", "Pendidikan", "SMA/SMK, D1, D3, D4, S1, S2, SMP, Lainnya"),
    ("fresh_grad", "Fresh Graduate", "1 = belum pernah bekerja"),
    ("exp_tele", "Pengalaman Tele (ada)", "Telemarketing/Telecollection/Deskcollection/Customer Care; 1 = pernah"),
    ("exp_tele_ord", "Pengalaman Tele (lama)", "0 tidak ada, 1 < 1 tahun, 2 >= 1 tahun, kosong = ada tapi lama tidak diketahui"),
    ("exp_finance", "Pengalaman Finance (ada)", "Multifinance/Perbankan; 1 = pernah"), ("exp_finance_ord", "Pengalaman Finance (lama)", "skala sama dengan Pengalaman Tele (lama)"),
    ("exp_sales", "Pengalaman Sales", "1 = pernah (tidak ditanyakan di form kuis -> kosong)"),
    ("jarak_km", "Jarak ke Cisauk (km)", "garis lurus; lihat Sumber Jarak untuk tingkat presisi"),
    ("jarak_src", "Sumber Jarak", "kecamatan (presisi) / kab/kota (kasar) / bracket (titik tengah pilihan jarak di form kuis) / tidak_terbaca"),
    ("lokasi_presisi", "Lokasi Presisi", "1 = jarak dari kecamatan"), ("jarak_bin", "Kelompok Jarak", "1-15, 16-30, 31-45, >45 km"),
    ("geo_match", "Domisili Terbaca", "hasil pembacaan domisili"),
    ("ext_jarak_bracket_km", "[Kuis] Jarak Pilihan Pelamar (km)", "hanya form kuis"), ("ext_ms_office", "[Kuis] MS Office", "0 tidak, 1 jarang, 2 sangat terbiasa"),
    ("ext_siap_join", "[Kuis] Siap Join", "1 = siap segera"),
    ("ext_menikah", "[Lama] Menikah", "hanya form lama"), ("ext_punya_anak", "[Lama] Punya Anak", "hanya form lama"),
    ("ext_nilai_norm", "[Lama] Nilai (0-100)", "hanya form lama"), ("ext_kuliah_sekarang", "[Lama] Sedang Kuliah", "hanya form lama"),
]


def to_excel(u: pd.DataFrame, path) -> pd.DataFrame:
    """Satu Excel lengkap & berlabel: 2024 + Des25/Jan26 (form lama) + semua Kuis 2026 + Data Diri 2026 (Tele Marketing)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    d = u.copy()
    d["kelompok"] = np.select(
        [(d["era"] == "v1") & (d["ts"] < "2025-01-01"), (d["era"] == "v1") & (d["ts"] >= "2025-12-01"), d["era"] == "v2", d["era"] == "v3"],
        [KELOMPOK["A"], KELOMPOK["B"], KELOMPOK["C"], KELOMPOK["D"]], default="")
    d = d[d["kelompok"] != ""]  # v1 selain Jun-Nov 2024 & Des 2025-Jan 2026: Rekap_All tidak memuatnya -> tanpa label, tidak dimasukkan
    d = d[~((d["era"] == "v1") & (d["ts"] >= "2024-12-01") & (d["ts"] < "2025-12-01"))]
    belum = d["label_note"].astype(str).str.startswith("belum diketahui")
    # submit terlalu baru (rekap belum lengkap): baris berlabel 1 pun TIDAK dipakai -> kalau dipakai, sampel hanya berisi yang lolos (bias)
    baru = d["ts"] > (d.loc[belum, "ts"].min() - pd.Timedelta(seconds=1)) if belum.any() else pd.Series(False, index=d.index)
    siap, alasan = [], []
    for (_, r), b in zip(d.iterrows(), baru):
        if pd.isna(r["y"]):
            siap.append("Tidak"); alasan.append(r["label_note"] if str(r["label_note"]).startswith("belum") else "label tidak jelas (" + str(r["label_note"]) + ")")
        elif b:
            siap.append("Tidak"); alasan.append("periode ini rekap belum lengkap: hanya pelamar yang sudah ada di rekap yang berlabel -> bias")
        elif r["era"] == "v3":
            siap.append("Hati-hati"); alasan.append("form ini kemungkinan diisi SETELAH kandidat dipanggil (95% 'lolos' pada Agustus) -> pola bisa bias")
        elif r["era"] == "v2" and r["ts"] < pd.Timestamp("2026-04-01"):
            siap.append("Hati-hati"); alasan.append("awal pemakaian form kuis (kemungkinan uji coba)")
        else:
            siap.append("Ya"); alasan.append("")
    d["siap"], d["siap_alasan"] = siap, alasan
    d = d.sort_values(["kelompok", "ts"]).reset_index(drop=True)
    d["id"] = range(1, len(d) + 1)
    out = d[[c for c, _n, _e in COLS]].copy()
    out.columns = [n for _c, n, _e in COLS]
    out["Tanggal Submit"] = pd.to_datetime(out["Tanggal Submit"]).dt.tz_localize(None)
    out["Tanggal Interview (Rekap)"] = pd.to_datetime(out["Tanggal Interview (Rekap)"], errors="coerce").dt.tz_localize(None)
    for c in out.columns:
        if out[c].dtype == float:
            out[c] = out[c].round(2)
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="DD/MM/YYYY HH:MM") as xw:
        out.to_excel(xw, sheet_name="Data_Latih", index=False)
        pd.DataFrame({"Kolom": [n for _c, n, _e in COLS], "Penjelasan": [e for _c, _n, e in COLS]}).to_excel(xw, sheet_name="Kamus_Kolom", index=False)
    wb = load_workbook(path)
    base = Font(name="Arial", size=10)
    head = PatternFill("solid", fgColor="1F3864")
    ws = wb["Data_Latih"]
    n = ws.max_row
    for row in ws.iter_rows():
        for c in row:
            c.font = base
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(name="Arial", size=10, bold=True, color="FFFFFF"), head, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    cols = list(out.columns)
    lab = cols.index("LABEL (1=lanjut interview)") + 1
    for r in range(2, n + 1):
        ws.cell(r, lab).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(r, lab).font = Font(name="Arial", size=10, bold=True)
    for i, c in enumerate(cols, 1):
        w = {"Nama": 26, "Keterangan Label": 34, "Alasan": 50, "Domisili Terbaca": 34, "Kelompok Data": 34, "NIK/KTP": 19}.get(c, 14)
        ws.column_dimensions[get_column_letter(i)].width = w
        if c.startswith("Tanggal"):
            for r in range(2, n + 1):
                ws.cell(r, i).number_format = "DD/MM/YYYY" if "Interview" in c else "DD/MM/YYYY HH:MM"
        if c in ("NIK/KTP", "No HP"):
            for r in range(2, n + 1):
                ws.cell(r, i).number_format = "@"
    kl = ws.parent["Kamus_Kolom"]
    for row in kl.iter_rows():
        for c in row:
            c.font = base
    for c in kl[1]:
        c.font, c.fill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), head
    kl.column_dimensions["A"].width, kl.column_dimensions["B"].width = 34, 110
    # ---- Ringkasan: rumus COUNTIFS ke Data_Latih (hidup: ikut berubah bila data diedit) ----
    rs = wb.create_sheet("Ringkasan")
    L = lambda name: get_column_letter(cols.index(name) + 1)
    kel, labc, siapc = L("Kelompok Data"), L("LABEL (1=lanjut interview)"), L("Siap Dipakai Training?")
    rng = lambda c: f"Data_Latih!${c}$2:${c}${n}"
    rs["A1"], rs["A1"].font = "Ringkasan dataset latih Tele Marketing", Font(name="Arial", size=12, bold=True)
    hdr = ["Kelompok Data", "Jumlah", "Label 1 (lanjut)", "Label 0", "Label kosong", "% lanjut (dari berlabel)", "Siap: Ya", "Siap: Hati-hati", "Siap: Tidak"]
    for j, h in enumerate(hdr, 1):
        c = rs.cell(3, j, h)
        c.font, c.fill, c.alignment = Font(name="Arial", size=10, bold=True, color="FFFFFF"), head, Alignment(wrap_text=True)
    names = list(KELOMPOK.values())
    for i, k in enumerate(names, 4):
        rs.cell(i, 1, k)
        rs.cell(i, 2, f'=COUNTIFS({rng(kel)},$A{i})')
        rs.cell(i, 3, f'=COUNTIFS({rng(kel)},$A{i},{rng(labc)},1)')
        rs.cell(i, 4, f'=COUNTIFS({rng(kel)},$A{i},{rng(labc)},0)')
        rs.cell(i, 5, f"=B{i}-C{i}-D{i}")
        rs.cell(i, 6, f"=IFERROR(C{i}/(C{i}+D{i}),0)")
        for j, v in zip((7, 8, 9), ("Ya", "Hati-hati", "Tidak")):
            rs.cell(i, j, f'=COUNTIFS({rng(kel)},$A{i},{rng(siapc)},"{v}")')
    t = 4 + len(names)
    rs.cell(t, 1, "Total")
    for j in (2, 3, 4, 5, 7, 8, 9):
        col = get_column_letter(j)
        rs.cell(t, j, f"=SUM({col}4:{col}{t-1})")
    rs.cell(t, 6, f"=IFERROR(C{t}/(C{t}+D{t}),0)")
    for r in range(4, t + 1):
        rs.cell(r, 6).number_format = "0%"
        for j in range(1, 10):
            rs.cell(r, j).font = Font(name="Arial", size=10, bold=(r == t))
    notes = ["Catatan:",
             "- Label 1 = pelamar ada di Rekap_All (unit Tele Marketing). Label 0 = tidak ada di rekap. Kosong = belum diketahui (submit terlalu baru / label tidak jelas).",
             "- Rekap_All hanya berisi interview sampai 02/09/2026; submit setelah ~19/08/2026 yang tidak ada di rekap diberi label KOSONG (bukan 0).",
             "- Des 2024 - Nov 2025 tidak dimasukkan: Rekap_All tidak memuat periode itu, jadi semua pelamarnya akan salah berlabel 0.",
             "- Untuk training gunakan baris 'Siap Dipakai Training?' = Ya (tambahkan 'Hati-hati' hanya setelah dicek). Lihat kolom Alasan.",
             "- Kolom IDENTITAS (Nama, No HP, NIK) hanya untuk pengecekan manual; HAPUS sebelum dibagikan / diunggah ke Colab."]
    for i, tx in enumerate(notes, t + 2):
        rs.cell(i, 1, tx).font = Font(name="Arial", size=10, bold=(i == t + 2))
    rs.column_dimensions["A"].width = 38
    for j in range(2, 10):
        rs.column_dimensions[get_column_letter(j)].width = 15
    rs.row_dimensions[3].height = 32
    wb.move_sheet("Ringkasan", offset=-2)
    wb.save(path)
    return out


def main():
    ap = argparse.ArgumentParser()
    for k in ("v1", "v2", "v3", "rekap"):
        ap.add_argument(f"--{k}", required=True)
    ap.add_argument("--out", default="training_data")
    a = ap.parse_args()
    v1, v2, v3 = (pd.read_excel(p) for p in (a.v1, a.v2, a.v3))
    u = mark_training(build_unified(v1, v2, v3, a.rekap), UNIFIED_WINDOWS)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    u.to_csv(out / "unified_all.csv", index=False)
    anon = u.drop(columns=["nama", "hp", "ktp", "geo_match", "person"]).assign(
        id=[f"P{i:04d}" for i in range(len(u))], bulan=u["ts"].dt.to_period("M").astype(str)).drop(columns=["ts", "rekap_date"])
    anon.to_csv(out / "unified_anonymized.csv", index=False)
    to_excel(u, out / "Data_Latih_Tele_Marketing.xlsx")
    (out / "laporan.md").write_text(report(u, evaluate(u)), encoding="utf-8")
    print(f"selesai: {len(u)} pelamar, {int(u['use_train'].sum())} dipakai training -> {out}/")


if __name__ == "__main__":
    main()
